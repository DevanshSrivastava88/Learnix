"""
Learnix feature matrix Excel.
Columns: Category | Feature/Command | Usage/Example | Test Status | Notes/Issues
Run any time to regenerate with latest test results.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv, os

# ── load latest test results from CSV (if available) ─────────────────────────
TEST_CSV = r"D:\Projects\learnix\test_casual_results.csv"
test_status = {}   # label -> status
test_reply  = {}   # label -> reply snippet

if os.path.exists(TEST_CSV):
    import re as _re
    with open(TEST_CSV, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            test_status[row["label"]] = row["status"]
            test_reply[row["label"]]  = row["reply"][:120] if row["reply"] else ""
    for full_label in list(test_status.keys()):
        prefix = next((tok for tok in full_label.split("_") if _re.match(r"^[A-Z]+[0-9]+$", tok)), None)
        if prefix and prefix != full_label:
            test_status[prefix] = test_status[full_label]
            test_reply[prefix]  = test_reply[full_label]

# ── feature data ──────────────────────────────────────────────────────────────
# (category, feature, usage, test_label, notes)
# test_label: maps to test CSV; "" = not directly tested; "~" = tested via chain
ROWS = [
    # ── CORE ──────────────────────────────────────────────────────────────────
    ("Core", "/start",   "First-time welcome + onboarding\nExample: /start",                           "",      ""),
    ("Core", "/help",    "Full command menu with all features\nExample: /help  or  'what can you do'", "CC1",   ""),
    ("Core", "/info",    "How the bot works explained plainly\nExample: /info  or  'explain yourself'","CC3",   ""),
    ("Core", "/cancel",  "Abort any active flow (goal creation, task wizard, etc.)\nExample: /cancel", "",      ""),
    ("Core", "/reset",   "Wipe all data and start fresh\nExample: /reset",                             "",      ""),

    # ── TASKS & HABITS ────────────────────────────────────────────────────────
    ("Tasks", "/newtask",            "Wizard to create a task step-by-step\nExample: /newtask → follow prompts",    "", ""),
    ("Tasks", "Create habit (natural)","Describe a recurring habit; bot creates it\nExample: 'i want to meditate every morning'\n'workout every day'","CC4",""),
    ("Tasks", "Create reminder (timed)","One-time reminder with a time\nExample: 'ping me to take my meds at 9pm'\n'remind me to call mom at 6pm'","CC6",""),
    ("Tasks", "Create reminder (no time)","Reminder without a time → stored as Unscheduled task in task list\nExample: 'don't let me forget to call dad'\n'remind me to buy groceries'","CC7","Fixed: no longer creates daily habit; stores as unscheduled one-time task"),
    ("Tasks", "Show task list",       "List all active tasks\nExample: 'what am i tracking'\n'give me my task list'\n/tasks","CC8",""),
    ("Tasks", "Mark done",            "Report a task complete\nExample: 'just knocked out my pushups'\n'done with reading'","CC10","Disambiguation when multiple similar tasks exist"),
    ("Tasks", "Skip today",           "Skip a habit and reschedule to tomorrow\nExample: 'gonna skip journaling for today'\n'nah skip the reading thing today'\n'skip it today'","CC11","Pronoun 'it/that' → asks which task (correct)"),
    ("Tasks", "Reschedule",           "Move a reminder to a new time\nExample: 'push my call dad reminder to 8pm'\n'move workout to 7am'","CC14","'Push my X reminder to Y' sometimes creates new reminder instead of rescheduling"),
    ("Tasks", "Mark important",       "Flag a task as high priority ⚡\nExample: 'my therapy appointment is really important, flag it'\n'mark workout as important'","CC13","Fixed: no longer blindly marks last-reminded task when no name given"),
    ("Tasks", "Delay/snooze",         "Snooze the most recent reminder\nExample: 'snooze the meds reminder by an hour'\n'delay 30 mins'","CC12",""),
    ("Tasks", "Break into steps",     "Decompose a task into sub-steps\nExample: 'help me break losing 10kg into steps'\n'break learning guitar into steps'","CC15",""),
    ("Tasks", "/deletetask",          "Delete a task by picking from list\nExample: /deletetask",    "",""),
    ("Tasks", "/pause + /resume",     "Pause/resume reminders for a task\nExample: /pause  or  'pause my workout'","",""),
    ("Tasks", "/complete",            "Mark a task permanently complete\nExample: /complete  or  'I finished guitar'","",""),

    # ── SCHEDULE ──────────────────────────────────────────────────────────────
    ("Schedule", "Show schedule",    "Full day view — automatics, habits with next-due time\nExample: 'give me a rundown of today'\n'am i forgetting anything today'\n/schedule","CC16",""),
    ("Schedule", "Set study time",   "Set daily study reminder time\nExample: 'i usually study around 10pm'\n/settime 09:00","CC22","Fixed: 'I usually study at X' now correctly sets study time via regex pre-check"),
    ("Schedule", "Set morning brief","Set morning summary time\nExample: 'i wake up at 7am, set my brief then'\n/setmorning 07:30","CC21",""),
    ("Schedule", "Set EOD time",     "Set end-of-day reflection time\nExample: /seteod 22:00  or  'set eod to 10pm'","",""),

    # ── ANALYTICS ─────────────────────────────────────────────────────────────
    ("Analytics", "Activity graph",  "Bar chart of daily completions\nExample: 'how productive have i been lately'\n/graph","CC18",""),
    ("Analytics", "Skip patterns",   "Chart of skip patterns by task/day\nExample: 'what do i skip the most'\n/skipgraph","CC19",""),

    # ── SETTINGS ──────────────────────────────────────────────────────────────
    ("Settings", "Show settings",    "View all current settings\nExample: 'what are my current settings'\n/settings","CC20",""),
    ("Settings", "Phone call alerts","Enable/disable missed-call Telegram notifications\nExample: 'let me know when someone calls me'\n/twilio on","CC23",""),
    ("Settings", "/clear",           "Delete all data permanently (asks confirmation)\nExample: /clear","",""),

    # ── STUDY ─────────────────────────────────────────────────────────────────
    ("Study", "Create goal",         "Create a learning goal with difficulty Easy/Medium/Hard\nExample: 'i want to get better at python'\n'i want to learn react'\n/goal","CC25",""),
    ("Study", "Show goals",          "List all study goals with status\nExample: 'what am i studying right now'\n'show my learning goals'\n/goals","CC26",""),
    ("Study", "Show topics",         "List topics under goals (✅🔄⏭⬜)\nExample: 'list my study topics'\n/topics","CC27",""),
    ("Study", "Add topic",           "Add a topic to an existing goal\nExample: /addtopic","",""),
    ("Study", "Start study session", "Bot quizzes you on next topic\nExample: 'quiz me on something'\n'let's study'\n/study","CC28",""),
    ("Study", "Study progress",      "Show overall progress % per goal\nExample: 'how far along am i in learning'\n/progress","CC29",""),
    ("Study", "Skip topic",          "Skip current topic, move to next\nExample: 'skip this topic'\n'next topic'","",""),
    ("Study", "/editgoal",           "Edit goal name/description/difficulty\nExample: /editgoal","",""),
    ("Study", "/deletegoal",         "Delete a goal and all its topics\nExample: /deletegoal","",""),
    ("Study", "/pausegoal",          "Pause/unpause a study goal\nExample: /pausegoal","",""),

    # ── CONTEXT MEMORY ────────────────────────────────────────────────────────
    ("Context", "Pronoun resolution","Bot resolves 'this'/'that'/'it' to last-mentioned item\nExample: 'remind me about this every morning'\n'ok remind me to do that every night at 10pm'","CTX2","Partial: works for explicit creates; fails for skip/reschedule via pronoun"),
    ("Context", "Chained reminders", "After discussing a topic, create reminder referencing it\nExample: [discuss sleep problems] → 'ok remind me to do that every night'","CTX15",""),
    ("Context", "Multi-turn topic",  "Carry subject across turns without repeating\nExample: [diet struggle] → 'remind me to track my calories every day'","CTX2","Works for task creation; context cleared by /cancel"),

    # ── VOICE ─────────────────────────────────────────────────────────────────
    ("Voice", "Voice notes",         "Send Telegram voice message — transcribed and processed exactly like text\nExample: say 'remind me to workout at 6pm'","","Not in automated test suite — requires manual voice send"),

    # ── MULTI-INTENT ──────────────────────────────────────────────────────────
    ("Advanced", "Multi-intent message","Bot handles compound instructions (takes first intent)\nExample: 'remind me to meditate at 7am and also track my water intake daily'","MX3","Only first task created; second clause dropped — known limitation of 8b model; timed part creates habit instead of reminder (known issue)"),
    ("Advanced", "Vague done",          "Bare confirmation asks for clarification\nExample: 'done!' → bot asks which task","MX2","Fixed: 'done!'/'Done!' now routes to done intent not chat"),
    ("Advanced", "Introspect",          "'what do you know about me' → shows settings/profile","MX5","LLM inconsistency on 8b model"),
]

# ── styles ────────────────────────────────────────────────────────────────────
STATUS_STYLES = {
    "PASS": ("C6EFCE", "375623"),   # green fill, green font
    "WARN": ("FFEB9C", "9C5700"),   # yellow fill, orange font
    "FAIL": ("FFC7CE", "9C0006"),   # red fill, red font
    "":    ("F5F5F5", "666666"),    # grey for untested
}

CAT_FILLS = {
    "Core":       "D9E1F2",
    "Tasks":      "E2EFDA",
    "Schedule":   "FFF2CC",
    "Analytics":  "FCE4D6",
    "Settings":   "EAD1DC",
    "Study":      "D0E4F1",
    "Context":    "E8D5F5",
    "Voice":      "F4CCFF",
    "Advanced":   "FFE4C4",
}

thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Learnix Features"

# ── headers ───────────────────────────────────────────────────────────────────
headers = ["Category", "Feature / Command", "Usage / Example", "Test Status", "Notes / Known Issues"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[1].height = 22

# ── rows ──────────────────────────────────────────────────────────────────────
for r, (cat, feat, usage, test_lbl, notes) in enumerate(ROWS, 2):
    cat_fill  = PatternFill("solid", fgColor=CAT_FILLS.get(cat, "FFFFFF"))
    data_fill = PatternFill("solid", fgColor="FAFAFA")

    # Col A — category
    c = ws.cell(row=r, column=1, value=cat)
    c.fill = cat_fill
    c.font = Font(bold=True, size=10)
    c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    c.border = border

    # Col B — feature name
    c = ws.cell(row=r, column=2, value=feat)
    c.fill = cat_fill
    c.font = Font(bold=True, size=10)
    c.alignment = WRAP
    c.border = border

    # Col C — usage
    c = ws.cell(row=r, column=3, value=usage)
    c.fill = data_fill
    c.font = Font(size=10)
    c.alignment = WRAP
    c.border = border

    # Col D — test status
    status = test_status.get(test_lbl, "") if test_lbl else ""
    reply  = test_reply.get(test_lbl, "")  if test_lbl else ""
    fill_hex, font_hex = STATUS_STYLES.get(status, STATUS_STYLES[""])
    status_label = status if status else ("(untested)" if not test_lbl else "—")
    if reply and status:
        status_label += f"\n\"{reply[:80]}...\""
    c = ws.cell(row=r, column=4, value=status_label)
    c.fill = PatternFill("solid", fgColor=fill_hex)
    c.font = Font(bold=(status in ("PASS","FAIL")), color=font_hex, size=9)
    c.alignment = WRAP
    c.border = border

    # Col E — notes
    c = ws.cell(row=r, column=5, value=notes if notes else "")
    c.fill = data_fill
    c.font = Font(size=9, italic=bool(notes))
    c.alignment = WRAP
    c.border = border

    # row height
    lines = max(usage.count("\n")+1, (status_label or "").count("\n")+1, (notes or "").count("\n")+1)
    ws.row_dimensions[r].height = max(30, lines * 15)

# ── summary row ───────────────────────────────────────────────────────────────
total   = len([x for x in ROWS if x[3]])
passed  = sum(1 for x in ROWS if test_status.get(x[3]) == "PASS")
warned  = sum(1 for x in ROWS if test_status.get(x[3]) == "WARN")
failed  = sum(1 for x in ROWS if test_status.get(x[3]) == "FAIL")

sr = len(ROWS) + 2
ws.cell(row=sr, column=1, value="SUMMARY").font = Font(bold=True)
ws.cell(row=sr, column=2, value=f"Tested: {total} features  |  PASS: {passed}  WARN: {warned}  FAIL: {failed}").font = Font(bold=True, color="1F3864")

# ── column widths ─────────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 52
ws.column_dimensions["D"].width = 22
ws.column_dimensions["E"].width = 42

ws.freeze_panes = "A2"

out = r"D:\Projects\learnix\learnix_features.xlsx"
wb.save(out)
print(f"Saved: {out}  (PASS:{passed} WARN:{warned} FAIL:{failed} / {total} tested)")
